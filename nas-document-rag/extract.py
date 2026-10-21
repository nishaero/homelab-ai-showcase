#!/usr/bin/env python3
"""Extract text from the read-only document cache into text/<sha1(path)>.json.

Reads only ~/nas-rag/docs (a copy). Never touches the NAS. Files with no text layer
(scans) are flagged needs_ocr=true for the OCR step. Resumable: skips existing output.
"""
import hashlib, json, subprocess, sys
from pathlib import Path

ROOT = Path.home() / "nas-rag"
DOCS, OUT = ROOT / "docs", ROOT / "text"
OUT.mkdir(exist_ok=True)


def pdf(p):
    r = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", str(p), "-"],
                       capture_output=True, timeout=180)
    return r.stdout.decode("utf-8", "replace")


def docx(p):
    import docx
    d = docx.Document(str(p))
    parts = [x.text for x in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def pptx(p):
    from pptx import Presentation
    out = []
    for i, s in enumerate(Presentation(str(p)).slides, 1):
        out.append(f"# slide {i}")
        for sh in s.shapes:
            if sh.has_text_frame:
                out.append(sh.text_frame.text)
    return "\n".join(out)


def odt(p):
    from odf import teletype, text as odtext
    from odf.opendocument import load
    d = load(str(p))
    return "\n".join(teletype.extractText(x) for x in d.getElementsByType(odtext.P))


def rtf(p):
    from striprtf.striprtf import rtf_to_text
    return rtf_to_text(p.read_text(errors="replace"))


def plain(p):
    return p.read_bytes().decode("utf-8", "replace")


HANDLERS = {".pdf": pdf, ".docx": docx, ".xlsx": xlsx, ".pptx": pptx, ".odt": odt,
            ".rtf": rtf, ".txt": plain, ".md": plain, ".csv": plain}


def main():
    rels = [l.rstrip("\n") for l in open(ROOT / "docs.lst", encoding="utf-8")]
    done = bad = skipped = 0
    for rel in rels:
        p = DOCS / rel
        out = OUT / (hashlib.sha1(rel.encode()).hexdigest() + ".json")
        if out.exists() or not p.exists():
            continue
        ext = p.suffix.lower()
        fn = HANDLERS.get(ext)
        rec = {"path": rel, "ext": ext, "size": p.stat().st_size, "text": "", "error": None}
        if fn is None:
            rec["error"] = "unsupported-format"  # e.g. legacy .doc: left for a converter step
            skipped += 1
        else:
            try:
                rec["text"] = fn(p).strip()
            except Exception as e:  # corrupt/encrypted file: record it, keep going
                rec["error"] = f"{type(e).__name__}: {e}"[:200]
                bad += 1
        rec["needs_ocr"] = ext == ".pdf" and len(rec["text"]) < 40 and rec["error"] is None
        out.write_text(json.dumps(rec, ensure_ascii=False))
        done += 1
        if done % 200 == 0:
            print(f"{done} extracted, {bad} errors, {skipped} unsupported", flush=True)
    print(f"finished: {done} new, {bad} errors, {skipped} unsupported", flush=True)


if __name__ == "__main__":
    sys.exit(main())
